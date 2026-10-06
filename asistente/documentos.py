"""Leer documentos (PDF, Word, texto, páginas web) y dividirlos en fragmentos."""

from __future__ import annotations

import io
import re
import urllib.request
from dataclasses import dataclass
from html.parser import HTMLParser

USER_AGENT = "Mozilla/5.0 (compatible; Jarvis-Estudio/1.0)"
EXTENSIONES = ["pdf", "docx", "txt", "md", "html", "htm", "csv"]


@dataclass
class Pagina:
    numero: int | None  # página del PDF (None si el formato no tiene páginas)
    texto: str


def leer_archivo(nombre: str, datos: bytes) -> list[Pagina]:
    ext = nombre.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return leer_pdf(datos)
    if ext == "docx":
        import docx
        doc = docx.Document(io.BytesIO(datos))
        partes = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                partes.append(" | ".join(c.text for c in row.cells))
        return [Pagina(None, "\n".join(partes))]
    texto = _decode(datos)
    if ext in {"html", "htm"}:
        texto = html_a_texto(texto)[1]
    return [Pagina(None, texto)]


def leer_pdf(datos: bytes) -> list[Pagina]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(datos))
    paginas = [Pagina(i + 1, (p.extract_text() or "").strip()) for i, p in enumerate(reader.pages)]
    if not any(p.texto for p in paginas):
        raise ValueError("El PDF no contiene texto seleccionable (¿es un escaneo?). "
                         "Pásalo antes por un programa de OCR.")
    return paginas


def _decode(datos: bytes) -> str:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return datos.decode(enc)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


class _TextoHTML(HTMLParser):
    SALTAR = {"script", "style", "noscript", "nav", "footer", "header", "svg", "form", "aside"}
    BLOQUE = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article"}

    def __init__(self):
        super().__init__()
        self.partes: list[str] = []
        self.titulo = ""
        self._saltando = 0
        self._en_titulo = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SALTAR:
            self._saltando += 1
        if tag == "title":
            self._en_titulo = True
        if tag in self.BLOQUE:
            self.partes.append("\n")
        if tag in {"h1", "h2", "h3"}:
            self.partes.append("## ")

    def handle_endtag(self, tag):
        if tag in self.SALTAR and self._saltando:
            self._saltando -= 1
        if tag == "title":
            self._en_titulo = False
        if tag in self.BLOQUE:
            self.partes.append("\n")

    def handle_data(self, data):
        if self._en_titulo:
            self.titulo += data
        elif not self._saltando:
            self.partes.append(data)


def html_a_texto(html: str) -> tuple[str, str]:
    """Devuelve (título, texto) de una página HTML."""
    parser = _TextoHTML()
    parser.feed(html)
    texto = re.sub(r"[ \t]+", " ", "".join(parser.partes))
    texto = re.sub(r"\n\s*\n+", "\n\n", texto).strip()
    return parser.titulo.strip(), texto


def descargar_url(url: str, timeout: int = 20) -> tuple[str, list[Pagina]]:
    """Descarga una página web o un PDF en línea. Devuelve (título, páginas)."""
    if not re.match(r"^https?://", url):
        raise ValueError("La dirección debe empezar por http:// o https://")
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        tipo = resp.headers.get("Content-Type", "")
        datos = resp.read(30_000_000)
    if "pdf" in tipo or url.lower().endswith(".pdf"):
        return url.rsplit("/", 1)[-1] or url, leer_pdf(datos)
    titulo, texto = html_a_texto(_decode(datos)) if "html" in tipo or not tipo else ("", _decode(datos))
    if not texto.strip():
        raise ValueError("No se encontró texto en la página.")
    return titulo or url, [Pagina(None, texto)]


def fragmentar(paginas: list[Pagina], tam: int = 1200, solape: int = 200) -> list[tuple[int | None, str]]:
    """Divide en fragmentos de ~`tam` caracteres, cortando por párrafos o frases.

    Devuelve una lista de (página, texto). El solape evita perder ideas que
    quedan partidas entre dos fragmentos.
    """
    fragmentos = []
    for pagina in paginas:
        texto = re.sub(r"[ \t]+", " ", pagina.texto).strip()
        inicio = 0
        while inicio < len(texto):
            fin = min(len(texto), inicio + tam)
            if fin < len(texto):
                corte = max(texto.rfind("\n\n", inicio, fin), texto.rfind(". ", inicio, fin))
                if corte > inicio + tam // 2:
                    fin = corte + 1
            trozo = texto[inicio:fin].strip()
            if trozo:
                fragmentos.append((pagina.numero, trozo))
            if fin >= len(texto):
                break
            inicio = max(fin - solape, inicio + 1)
    return fragmentos
