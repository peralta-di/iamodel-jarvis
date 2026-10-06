"""Convierte los resultados en archivos: Word, PowerPoint, diagramas y tarjetas para Anki."""

from __future__ import annotations

import csv
import html
import io
import json
import re

COLORES = ["#4C6EF5", "#12B886", "#F59F00", "#E64980", "#7950F2", "#15AABF", "#FA5252", "#82C91E"]


# --------------------------------------------------------------- mapa mental
def _q(texto: str) -> str:
    return '"' + texto.replace("\\", "\\\\").replace('"', '\\"') + '"'


def mapa_a_dot(mapa: dict) -> str:
    """Mapa mental → Graphviz DOT (se dibuja en la app y se puede abrir en cualquier visor)."""
    lineas = ["digraph mapa {", "  graph [rankdir=LR, bgcolor=transparent, nodesep=0.25, ranksep=0.6];",
              '  node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=12, '
              'color="#00000000"];', '  edge [color="#888888", arrowhead=none];',
              f'  raiz [label={_q(mapa["titulo"])}, fillcolor="#212529", fontcolor=white, fontsize=16];']
    for i, rama in enumerate(mapa.get("ramas", [])):
        color = COLORES[i % len(COLORES)]
        rid = f"r{i}"
        lineas.append(f'  {rid} [label={_q(rama["titulo"])}, fillcolor="{color}", fontcolor=white];')
        lineas.append(f"  raiz -> {rid} [color=\"{color}\", penwidth=2];")
        for j, sub in enumerate(rama.get("subramas", [])):
            sid = f"{rid}s{j}"
            lineas.append(f'  {sid} [label={_q(sub["titulo"])}, fillcolor="{color}33"];')
            lineas.append(f"  {rid} -> {sid} [color=\"{color}\"];")
            for k, det in enumerate(sub.get("detalles", [])):
                did = f"{sid}d{k}"
                lineas.append(f'  {did} [label={_q(det)}, shape=plaintext, style="", fontsize=10];')
                lineas.append(f"  {sid} -> {did} [color=\"{color}88\"];")
    lineas.append("}")
    return "\n".join(lineas)


def _mermaid_texto(texto: str) -> str:
    return re.sub(r"[()\[\]{}]", " ", texto).strip() or "-"


def mapa_a_mermaid(mapa: dict) -> str:
    lineas = ["mindmap", f"  root(({_mermaid_texto(mapa['titulo'])}))"]
    for rama in mapa.get("ramas", []):
        lineas.append(f"    {_mermaid_texto(rama['titulo'])}")
        for sub in rama.get("subramas", []):
            lineas.append(f"      {_mermaid_texto(sub['titulo'])}")
            for det in sub.get("detalles", []):
                lineas.append(f"        {_mermaid_texto(det)}")
    return "\n".join(lineas)


def mapa_a_markdown(mapa: dict) -> str:
    lineas = [f"# {mapa['titulo']}"]
    for rama in mapa.get("ramas", []):
        lineas.append(f"\n## {rama['titulo']}")
        for sub in rama.get("subramas", []):
            lineas.append(f"- **{sub['titulo']}**")
            lineas += [f"  - {d}" for d in sub.get("detalles", [])]
    return "\n".join(lineas)


def mermaid_a_html(codigo: str, titulo: str) -> str:
    """Página HTML independiente que dibuja el diagrama (abre en el navegador e imprime a PDF)."""
    return f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><title>{html.escape(titulo)}</title>
<style>body{{font-family:Helvetica,Arial,sans-serif;margin:24px;background:#fff}}</style></head>
<body><h2>{html.escape(titulo)}</h2>
<pre class="mermaid">
{html.escape(codigo)}
</pre>
<script type="module">
import mermaid from "https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs";
mermaid.initialize({{startOnLoad: true, theme: "default"}});
</script></body></html>"""


# ----------------------------------------------------------------- diagramas
FORMAS_DOT = {"caja": "box", "rombo": "diamond", "ovalo": "ellipse", "circulo": "circle"}
FORMAS_MERMAID = {"caja": ("[", "]"), "rombo": ("{", "}"), "ovalo": ("([", "])"), "circulo": ("((", "))")}


def diagrama_a_dot(diagrama: dict, tipo: str = "") -> str:
    direccion = "LR" if tipo in {"Mapa conceptual", "Ciclo"} else "TB"
    lineas = ["digraph diagrama {", f"  graph [rankdir={direccion}, bgcolor=transparent];",
              '  node [style=filled, fillcolor="#E7F5FF", color="#1C7ED6", fontname="Helvetica"];',
              '  edge [color="#495057", fontname="Helvetica", fontsize=10];']
    ids = set()
    for n in diagrama.get("nodos", []):
        ids.add(n["id"])
        forma = FORMAS_DOT.get(n.get("forma", "caja"), "box")
        estilo = ', style="rounded,filled"' if forma == "box" else ""
        lineas.append(f'  {_q(n["id"])} [label={_q(n["texto"])}, shape={forma}{estilo}];')
    for c in diagrama.get("conexiones", []):
        if c["desde"] in ids and c["hasta"] in ids:
            etiqueta = f" [label={_q(c['etiqueta'])}]" if c.get("etiqueta") else ""
            lineas.append(f"  {_q(c['desde'])} -> {_q(c['hasta'])}{etiqueta};")
    lineas.append("}")
    return "\n".join(lineas)


def diagrama_a_mermaid(diagrama: dict, tipo: str = "") -> str:
    direccion = "LR" if tipo in {"Mapa conceptual", "Ciclo"} else "TD"
    seguro = lambda s: re.sub(r"\W", "_", s)
    lineas = [f"flowchart {direccion}"]
    for n in diagrama.get("nodos", []):
        a, b = FORMAS_MERMAID.get(n.get("forma", "caja"), ("[", "]"))
        lineas.append(f'  {seguro(n["id"])}{a}"{n["texto"].replace(chr(34), "")}"{b}')
    for c in diagrama.get("conexiones", []):
        etiqueta = f'|"{c["etiqueta"]}"|' if c.get("etiqueta") else ""
        lineas.append(f"  {seguro(c['desde'])} -->{etiqueta} {seguro(c['hasta'])}")
    return "\n".join(lineas)


# -------------------------------------------------------------- PowerPoint
def presentacion_a_pptx(pres: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    acento = RGBColor(0x4C, 0x6E, 0xF5)

    portada = prs.slides.add_slide(prs.slide_layouts[0])
    portada.shapes.title.text = pres.get("titulo", "Presentación")
    portada.placeholders[1].text = pres.get("subtitulo", "")
    portada.shapes.title.text_frame.paragraphs[0].font.color.rgb = acento

    for d in pres.get("diapositivas", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = d.get("titulo", "")
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = acento
        cuerpo = slide.placeholders[1]
        cuerpo.left, cuerpo.top, cuerpo.width, cuerpo.height = Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.3)
        tf = cuerpo.text_frame
        for i, punto in enumerate(d.get("puntos", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = punto
            p.font.size = Pt(24)
        if d.get("notas"):
            slide.notes_slide.notes_text_frame.text = d["notas"]

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def presentacion_a_markdown(pres: dict) -> str:
    partes = [f"# {pres.get('titulo', '')}\n\n_{pres.get('subtitulo', '')}_"]
    for n, d in enumerate(pres.get("diapositivas", []), 1):
        puntos = "\n".join(f"- {p}" for p in d.get("puntos", []))
        partes.append(f"## {n}. {d.get('titulo', '')}\n\n{puntos}\n\n> Notas: {d.get('notas', '')}")
    return "\n\n".join(partes)


# --------------------------------------------------------------------- Word
def markdown_a_docx(markdown: str, titulo: str = "") -> bytes:
    """Conversión sencilla de Markdown (títulos, listas, tablas, negritas) a Word."""
    import docx

    doc = docx.Document()
    if titulo:
        doc.add_heading(titulo, level=0)
    lineas = markdown.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        if linea.startswith("|") and i + 1 < len(lineas) and re.match(r"^\|[\s:|-]+\|$", lineas[i + 1].strip()):
            filas = [linea]
            i += 2
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i].strip())
                i += 1
            celdas = [[c.strip() for c in f.strip("|").split("|")] for f in filas]
            ancho = max(len(f) for f in celdas)
            tabla = doc.add_table(rows=len(celdas), cols=ancho)
            tabla.style = "Table Grid"
            for r, fila in enumerate(celdas):
                for c, valor in enumerate(fila):
                    _texto_con_formato(tabla.cell(r, c).paragraphs[0], valor, negrita=r == 0)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", linea)
        if m:
            doc.add_heading(_limpiar(m.group(2)), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*+]\s+", linea):
            nivel = (len(linea) - len(linea.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if nivel else "List Bullet")
            _texto_con_formato(p, re.sub(r"^\s*[-*+]\s+", "", linea))
        elif re.match(r"^\s*\d+[.)]\s+", linea):
            p = doc.add_paragraph(style="List Number")
            _texto_con_formato(p, re.sub(r"^\s*\d+[.)]\s+", "", linea))
        elif linea.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _texto_con_formato(p, linea.lstrip("> "))
        elif linea.strip() and not re.match(r"^-{3,}$", linea.strip()):
            _texto_con_formato(doc.add_paragraph(), linea)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _sin_html(texto: str) -> str:
    return re.sub(r"</?(details|summary)>", " ", texto).replace("`", "")


def _limpiar(texto: str) -> str:
    return _sin_html(texto).replace("**", "")


def _texto_con_formato(parrafo, texto: str, negrita: bool = False) -> None:
    for i, trozo in enumerate(re.split(r"\*\*(.+?)\*\*", _sin_html(texto))):
        if trozo:
            parrafo.add_run(trozo).bold = negrita or i % 2 == 1


# ------------------------------------------------------------- tarjetas / quiz
def tarjetas_a_csv(tarjetas: list[dict]) -> bytes:
    """CSV importable en Anki (Archivo → Importar), Quizlet, etc."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    for t in tarjetas:
        writer.writerow([t["frente"], t["reverso"]])
    return buf.getvalue().encode("utf-8")


def quiz_a_markdown(quiz: dict, con_respuestas: bool = True) -> str:
    partes = [f"# {quiz.get('titulo', 'Quiz')}"]
    for n, p in enumerate(quiz.get("preguntas", []), 1):
        bloque = f"**{n}. {p['pregunta']}**"
        if p.get("opciones"):
            bloque += "\n" + "\n".join(f"   {chr(97 + i)}) {o}" for i, o in enumerate(p["opciones"]))
        partes.append(bloque)
    if con_respuestas:
        partes.append("---\n## Respuestas")
        for n, p in enumerate(quiz.get("preguntas", []), 1):
            partes.append(f"{n}. **{p['respuesta_correcta']}** — {p['explicacion']}")
    return "\n\n".join(partes)


def a_json(datos) -> bytes:
    return json.dumps(datos, ensure_ascii=False, indent=2).encode("utf-8")
