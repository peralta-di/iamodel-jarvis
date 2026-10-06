import io
import json
from types import SimpleNamespace

import pytest

from asistente import ciencia, estudio, exportar
from asistente.almacen import Almacen
from asistente.busqueda import BM25
from asistente.documentos import Pagina, fragmentar, html_a_texto, leer_archivo
from asistente.herramientas import Herramientas
from asistente.motores import ClaudeMotor, Motor, Respuesta, _a_texto_plano


# ------------------------------------------------------------------ utilidades
def pdf_minimo(paginas: list[str]) -> bytes:
    """Construye un PDF válido con una línea de texto por página."""
    objetos = ["<< /Type /Catalog /Pages 2 0 R >>", None,
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for texto in paginas:
        stream = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET"
        objetos.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        contenido = len(objetos)
        objetos.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {contenido} 0 R >>")
        kids.append(f"{len(objetos)} 0 R")
    objetos[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objetos, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


class MotorFalso(Motor):
    """Motor de prueba: devuelve respuestas fijas y guarda lo que recibe."""
    nombre = "falso"

    def __init__(self, json_respuestas=None, max_contexto_chars=20_000):
        self.llamadas = []
        self.json_respuestas = list(json_respuestas or [])
        self.max_contexto_chars = max_contexto_chars

    def conversar(self, sistema, historial, herramientas=None, web=False, al_progreso=None):
        self.llamadas.append({"sistema": sistema, "historial": historial, "web": web})
        texto = f"respuesta {len(self.llamadas)}"
        return Respuesta(texto=texto, api=[{"role": "assistant", "content": texto}])

    def completar_json(self, prompt, esquema, sistema=""):
        self.llamadas.append({"prompt": prompt})
        return self.json_respuestas.pop(0)


# --------------------------------------------------------------------- ciencia
def test_calcular_operaciones():
    assert ciencia.calcular("resolver", "x^2 - 5x + 6 = 0")["resultado"] == "[{x: 2}, {x: 3}]"
    assert ciencia.calcular("resolver", "x + y = 3, x - y = 1")["resultado"] == "[{x: 2, y: 1}]"
    assert ciencia.calcular("integrar", "x^2", "x", "0", "3")["resultado"] == "9"
    assert ciencia.calcular("derivar", "x^3", "x", orden=2)["resultado"] == "6*x"
    assert ciencia.calcular("limite", "sin(x)/x", "x", hasta="0")["resultado"] == "1"
    assert ciencia.calcular("matriz", "[[1,2],[3,4]]")["determinante"] == "-2"
    assert ciencia.calcular("evaluar", "sqrt(2)")["aproximado"].startswith("1.41421")


def test_calcular_rechaza_codigo():
    with pytest.raises(ValueError):
        ciencia.calcular("evaluar", "__import__('os').system('ls')")
    with pytest.raises(ValueError):
        ciencia.calcular("nada", "x")


def test_quimica():
    assert ciencia.masa_molar("H2O")["masa_molar_g_mol"] == pytest.approx(18.015, abs=0.01)
    assert ciencia.masa_molar("Ca(OH)2")["masa_molar_g_mol"] == pytest.approx(74.09, abs=0.01)
    assert ciencia.balancear_ecuacion("Fe + O2 -> Fe2O3")["ecuacion_ajustada"] == "4 Fe + 3 O2 -> 2 Fe2O3"
    with pytest.raises(ValueError):
        ciencia.masa_molar("Xx2")


def test_herramientas_devuelven_errores_al_modelo():
    h = Herramientas()
    texto, error = h.ejecutar("calcular", {"operacion": "resolver", "expresion": "2x = 4"})
    assert not error and "2" in json.loads(texto)["resultado"]
    texto, error = h.ejecutar("masa_molar", {"formula": "Qq"})
    assert error and texto.startswith("Error")
    assert [d["name"] for d in h.definiciones()] == ["calcular", "masa_molar", "balancear_ecuacion"]


# ------------------------------------------------------------------ documentos
def test_leer_pdf_y_html():
    paginas = leer_archivo("apuntes.pdf", pdf_minimo(["Hola fotosintesis", "Segunda pagina"]))
    assert [p.numero for p in paginas] == [1, 2]
    assert "fotosintesis" in paginas[0].texto
    titulo, texto = html_a_texto("<html><title>T</title><script>x=1</script><body><h1>Hola</h1>"
                                 "<p>Mundo</p></body></html>")
    assert titulo == "T" and "Hola" in texto and "Mundo" in texto and "x=1" not in texto


def test_fragmentar_con_solape():
    texto = " ".join(f"Frase número {i}." for i in range(400))
    trozos = fragmentar([Pagina(3, texto)], tam=500, solape=100)
    assert len(trozos) > 5
    assert all(pag == 3 and len(t) <= 500 for pag, t in trozos)


def test_bm25_encuentra_el_fragmento_relevante():
    idx = BM25([
        {"id": 1, "texto": "La fotosíntesis ocurre en los cloroplastos de las plantas.", "documento": "bio", "pagina": 1},
        {"id": 2, "texto": "La segunda ley de Newton relaciona fuerza, masa y aceleración.", "documento": "fis", "pagina": 2},
        {"id": 3, "texto": "Los ácidos tienen un pH menor que siete.", "documento": "qui", "pagina": None},
    ])
    assert idx.buscar("¿Qué dice la ley de Newton sobre la fuerza?")[0].id == 2
    assert idx.buscar("fotosintesis plantas")[0].documento == "bio"
    assert idx.buscar("xyz") == []


# --------------------------------------------------------------------- almacén
def test_almacen_chats_y_documentos():
    db = Almacen(":memory:")
    chat = db.crear_chat()
    texto = " ".join(f"Oración {i} sobre biología celular." for i in range(300))
    doc = db.agregar_documento(chat, "bio.pdf", [Pagina(1, texto), Pagina(2, "Mitocondria.")])
    assert db.documentos(chat)[0]["paginas"] == 2
    completo = db.texto_documento(doc)
    assert completo.count("Oración 150 ") == 1  # el solape se elimina
    assert "[Página 2]" in completo
    assert db.indice(chat).buscar("mitocondria")[0].pagina == 2
    db.borrar_chat(chat)
    assert db.listar_chats() == []


# ----------------------------------------------------------------- motores
def test_a_texto_plano_quita_bloques_internos():
    historial = [
        {"role": "user", "content": [{"type": "image", "source": {"type": "base64", "data": "AAA"}},
                                     {"type": "text", "text": "¿Qué es?"}]},
        {"role": "assistant", "content": [{"type": "thinking", "thinking": ""},
                                          {"type": "tool_use", "id": "1", "name": "calcular", "input": {}}]},
        {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "1", "content": "9"}]},
        {"role": "assistant", "content": [{"type": "text", "text": "Es 9."}]},
    ]
    plano = _a_texto_plano(historial)
    assert plano == [{"role": "user", "content": "¿Qué es?", "images": ["AAA"]},
                     {"role": "assistant", "content": "Es 9."}]


def _bloque(**campos):
    return SimpleNamespace(**campos, to_dict=lambda: {k: v for k, v in campos.items()})


class ClienteFalso:
    """Imita client.beta.messages.create con respuestas preparadas."""

    def __init__(self, respuestas):
        self.respuestas = list(respuestas)
        self.peticiones = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))

    def _create(self, **kwargs):
        self.peticiones.append(json.loads(json.dumps(kwargs, default=str)))
        return self.respuestas.pop(0)


def test_claude_bucle_de_herramientas():
    cliente = ClienteFalso([
        SimpleNamespace(stop_reason="tool_use", content=[
            _bloque(type="tool_use", id="t1", name="calcular",
                    input={"operacion": "integrar", "expresion": "x^2", "variable": "x", "desde": "0", "hasta": "3"})]),
        SimpleNamespace(stop_reason="end_turn", content=[
            _bloque(type="text", text="La integral vale 9.", citations=None)]),
    ])
    motor = ClaudeMotor(cliente=cliente)
    r = motor.conversar("sistema", [{"role": "user", "content": "integra x^2 de 0 a 3"}], Herramientas())
    assert r.texto == "La integral vale 9."
    assert [m["role"] for m in r.api] == ["assistant", "user", "assistant"]
    resultado = r.api[1]["content"][0]
    assert resultado["tool_use_id"] == "t1" and json.loads(resultado["content"])["resultado"] == "9"
    peticion = cliente.peticiones[1]
    assert peticion["model"] == "claude-opus-5-5" and peticion["thinking"] == {"type": "adaptive"}
    assert peticion["fallbacks"] == "default"
    assert len(peticion["messages"]) == 3  # pregunta + tool_use + tool_result
    assert r.herramientas[0]["nombre"] == "calcular"


def test_claude_recoge_fuentes_web():
    resultados = [SimpleNamespace(url="https://a.org", title="A"), SimpleNamespace(url="https://a.org", title="A")]
    cliente = ClienteFalso([SimpleNamespace(stop_reason="end_turn", content=[
        _bloque(type="web_search_tool_result", content=resultados),
        _bloque(type="text", text="Dato.", citations=[SimpleNamespace(url="https://b.org", title="B")])])])
    r = ClaudeMotor(cliente=cliente).conversar("", [{"role": "user", "content": "hola"}], web=True)
    assert r.fuentes == [{"titulo": "A", "url": "https://a.org"}, {"titulo": "B", "url": "https://b.org"}]
    assert {t["name"] for t in cliente.peticiones[0]["tools"]} == {"web_search", "web_fetch"}


# --------------------------------------------------------------------- estudio
def test_responder_usa_documentos_y_guarda_historial():
    db = Almacen(":memory:")
    chat = db.crear_chat()
    db.agregar_documento(chat, "fisica.pdf", [Pagina(4, "La segunda ley de Newton dice que F = m a.")])
    motor = MotorFalso()
    estudio.responder(motor, db, chat, "¿Qué dice la ley de Newton?", solo_documentos=True, web=True)
    llamada = motor.llamadas[0]
    pregunta = llamada["historial"][-1]["content"][-1]["text"]
    assert "[1] (fisica.pdf, pág. 4)" in pregunta and "SOLO MIS DOCUMENTOS" in llamada["sistema"]
    assert llamada["web"] is False  # el modo solo documentos desactiva la web

    estudio.responder(motor, db, chat, "¿Y la tercera?")
    historial = motor.llamadas[1]["historial"]
    assert [m["role"] for m in historial] == ["user", "assistant", "user"]
    mensajes = db.mensajes(chat)
    assert mensajes[0]["texto"] == "¿Qué dice la ley de Newton?"  # se muestra sin el contexto
    assert mensajes[1]["extra"]["documentos"][0]["pagina"] == 4
    assert db.listar_chats()[0]["titulo"] == "¿Qué dice la ley de Newton?"


def test_resumir_por_partes_si_el_texto_es_largo():
    motor = MotorFalso(max_contexto_chars=1000)
    salida = estudio.resumir(motor, "palabra " * 500, "Puntos clave")
    assert len(motor.llamadas) == 6  # 5 partes + combinación final
    assert salida == "respuesta 6"


def test_investigar_planifica_e_informa():
    motor = MotorFalso(json_respuestas=[{"subpreguntas": ["A", "B", "C"]}])
    r = estudio.investigar(motor, "tema", subpreguntas=2, web=True)
    assert r["subpreguntas"] == ["A", "B"]
    assert len(r["hallazgos"]) == 2
    assert "no puede buscar en internet" in r["informe"]


# -------------------------------------------------------------------- exportar
MAPA = {"titulo": "Célula", "ramas": [{"titulo": "Partes", "subramas": [
    {"titulo": "Núcleo", "detalles": ["Contiene \"ADN\""]}]}]}


def test_exportar_mapa_y_diagrama():
    dot = exportar.mapa_a_dot(MAPA)
    assert dot.startswith("digraph") and '\\"ADN\\"' in dot
    assert exportar.mapa_a_mermaid(MAPA).startswith("mindmap")
    diagrama = {"titulo": "t", "nodos": [{"id": "n1", "texto": "Inicio", "forma": "ovalo"},
                                         {"id": "n2", "texto": "¿x>0?", "forma": "rombo"}],
                "conexiones": [{"desde": "n1", "hasta": "n2", "etiqueta": ""},
                               {"desde": "n2", "hasta": "nX", "etiqueta": "Sí"}]}
    dot = exportar.diagrama_a_dot(diagrama)
    assert "shape=diamond" in dot and "nX" not in dot
    assert "-->" in exportar.diagrama_a_mermaid(diagrama)


def test_exportar_pptx_y_docx():
    from pptx import Presentation
    import docx
    pres = {"titulo": "Física", "subtitulo": "Tema 1", "diapositivas": [
        {"titulo": "Newton", "puntos": ["F = m a", "Acción y reacción"], "notas": "Explicar"}]}
    prs = Presentation(io.BytesIO(exportar.presentacion_a_pptx(pres)))
    assert len(prs.slides) == 2
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Explicar"
    md = "# Título\n\n- **uno**\n- dos\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\nTexto final."
    d = docx.Document(io.BytesIO(exportar.markdown_a_docx(md)))
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "2"
    assert exportar.tarjetas_a_csv([{"frente": "a,b", "reverso": "c"}]).decode() == '"a,b",c\r\n'
